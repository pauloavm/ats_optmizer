import streamlit as st
from openai import OpenAI
import os
from dotenv import load_dotenv
import json
import re
import PyPDF2
import docx
from fpdf import FPDF, XPos, YPos
import io

# 1. Carrega as variaveis do arquivo .env (Seguranca)
load_dotenv()
api_key = os.getenv("OPENAI_API_KEY")
client = OpenAI(api_key=api_key)

# 2. Configuracao da Pagina Streamlit
st.set_page_config(page_title="ATS Optimizer Pro", layout="wide")

# 3. Funcoes de Extracao e Pre-processamento (Camada Silver - Zero IA)
def ler_docx(arquivo):
    doc = docx.Document(arquivo)
    return "\n".join([paragrafo.text for paragrafo in doc.paragraphs])

def converter_pdf_para_otimizado(arquivo_pdf):
    leitor = PyPDF2.PdfReader(arquivo_pdf)
    texto_completo = ""
    for pagina in leitor.pages:
        texto_completo += pagina.extract_text() + "\n"

    linhas = texto_completo.split('\n')
    cv_estruturado = {"cabecalho": [], "secoes": {}}
    secao_atual = "cabecalho"

    keywords_secoes = [
        'EXPERIÊNCIA', 'EXPERIÊNCIAS', 'EDUCAÇÃO', 'FORMAÇÃO', 
        'RESUMO', 'IDIOMAS', 'COMPETÊNCIAS', 'HABILIDADES', 'CONQUISTAS',
        'CERTIFICATIONS', 'LANGUAGES'
    ]

    for linha in linhas:
        linha_limpa = linha.strip()
        if not linha_limpa:
            continue

        linha_upper = linha_limpa.upper()
        is_secao = any(linha_upper == kw or linha_upper.startswith(kw) for kw in keywords_secoes)

        if is_secao and len(linha_limpa) < 35:
            secao_atual = linha_upper
            cv_estruturado["secoes"][secao_atual] = []
        else:
            if secao_atual == "cabecalho":
                cv_estruturado["cabecalho"].append(linha_limpa)
            else:
                if secao_atual not in cv_estruturado["secoes"]:
                    cv_estruturado["secoes"][secao_atual] = []
                cv_estruturado["secoes"][secao_atual].append(linha_limpa)

    md_linhas = []
    if cv_estruturado["cabecalho"]:
        md_linhas.append("### INFORMAÇÕES PESSOAIS")
        md_linhas.extend(cv_estruturado["cabecalho"])
        md_linhas.append("")

    for secao, conteudo in cv_estruturado["secoes"].items():
        md_linhas.append(f"### {secao}")
        for item in conteudo:
            if item.startswith('•') or item.startswith('-') or item.startswith('▪'):
                md_linhas.append(f"- {item[1:].strip()}")
            else:
                item_limpo = re.sub(r'\s+', ' ', item)
                md_linhas.append(item_limpo)
        md_linhas.append("")

    return "\n".join(md_linhas)

def limpar_texto_para_tokens(texto):
    texto_limpo = re.sub(r'\s+', ' ', texto)
    return texto_limpo.strip()

# Funcoes de Exportacao (PDF e DOCX)
def gerar_pdf(texto_md):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=11)

    for linha in texto_md.split('\n'):
        # Limpeza basica do markdown para o PDF
        linha_limpa = linha.replace('### ', '').replace('**', '').strip()
        # Tratamento de encoding para evitar erros com acentos no FPDF padrao
        linha_limpa = linha_limpa.encode('latin-1', 'replace').decode('latin-1')

        if linha_limpa:
            # Sintaxe atualizada do fpdf2 usando XPos e YPos
            pdf.cell(0, 6, txt=linha_limpa, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        else:
            pdf.cell(0, 4, txt="", new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    # No fpdf2, output() sem argumentos retorna um bytearray, perfeito para o Streamlit
    return pdf.output()

def gerar_docx(texto_md):
    doc = docx.Document()
    for linha in texto_md.split('\n'):
        linha_limpa = linha.replace('### ', '').replace('**', '').strip()
        if linha_limpa:
            doc.add_paragraph(linha_limpa)

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

# 4. Interface do Usuario (Camada Gold)
st.title("🎯 ATS Optimizer Pro: Análise e Conversão")
st.markdown("Diagnóstico profundo de aderência e otimização focada nos algoritmos da Gupy, inHire, LinkedIn e outros.")

col1, col2 = st.columns(2)
with col1:
    vaga_texto = st.text_area("Cole a descrição da vaga aqui", height=250)
with col2:
    arquivo_cv = st.file_uploader("Envie seu CV do LinkedIn (Profile.pdf)", type=["pdf", "docx"])

# 5. Motor de IA e Processamento
if st.button("Analisar Aderência e Gerar CV"):
    if not api_key:
        st.error("Erro: Chave da OpenAI não encontrada no arquivo .env.")
    elif vaga_texto and arquivo_cv:
        with st.spinner("Analisando algoritmos de ATS e processando com gpt-4o-mini..."):

            if arquivo_cv.name.endswith('.pdf'):
                cv_texto_bruto = converter_pdf_para_otimizado(arquivo_cv)
            else:
                cv_texto_bruto = ler_docx(arquivo_cv)

            vaga_otimizada = limpar_texto_para_tokens(vaga_texto)
            cv_otimizado = limpar_texto_para_tokens(cv_texto_bruto)

            prompt_sistema = """
            Você é um especialista em ATS e recrutamento tech no Brasil. 
            Sua missão é analisar o CV do candidato em relação à vaga e gerar um JSON estruturado com a análise, dicas por plataforma e o CV otimizado.
            NÃO INVENTE DADOS.

            Retorne EXATAMENTE no formato JSON abaixo:
            {
              "analise": {
                "score_estimado": "X%",
                "pontos_conexao": ["Força 1 encontrada no CV que bate com a vaga", "Força 2"],
                "pontos_melhoria": ["Gap 1 que o candidato precisa desenvolver ou destacar melhor", "Gap 2"]
              },
              "dicas_plataformas": {
                "Gupy": "Dica específica de como preencher os campos na Gupy baseada nesta vaga.",
                "inHire": "Dica específica para a inHire.",
                "Catho": "Dica específica para a Catho.",
                "Indeed": "Dica específica para o Indeed.",
                "Infojobs": "Dica específica para o Infojobs.",
                "LinkedIn": "Dica de SEO e palavras-chave para o perfil do LinkedIn."
              },
              "cv_adaptado_markdown": "Texto do CV otimizado em Markdown, seguindo a estrutura: NOME, Título, Contatos, SUMÁRIO, SISTEMAS, SKILLS, COMPETÊNCIAS, REALIZAÇÕES, EDUCAÇÃO, EXPERIÊNCIAS PROFISSIONAIS."
            }
            """

            try:
                resposta = client.chat.completions.create(
                    model="gpt-4o-mini",
                    response_format={ "type": "json_object" },
                    messages=[
                        {"role": "system", "content": prompt_sistema},
                        {"role": "user", "content": f"VAGA: {vaga_otimizada}\n\nCV: {cv_otimizado}"}
                    ],
                    temperature=0.1
                )

                resultado_json = json.loads(resposta.choices[0].message.content)

                # Exibicao da Analise
                st.header(f"📊 Score de Aderência: {resultado_json['analise']['score_estimado']}")

                col_analise1, col_analise2 = st.columns(2)
                with col_analise1:
                    st.subheader("✅ Pontos de Conexão (Suas Forças)")
                    for ponto in resultado_json['analise']['pontos_conexao']:
                        st.write(f"- {ponto}")

                with col_analise2:
                    st.subheader("⚠️ Pontos de Melhoria (Gaps)")
                    for gap in resultado_json['analise']['pontos_melhoria']:
                        st.write(f"- {gap}")

                st.divider()

                # Exibicao das Dicas por ATS
                st.header("🤖 Estratégia de Preenchimento por ATS")
                st.markdown("Como adaptar seu perfil para os diferentes algoritmos do mercado:")

                dicas = resultado_json['dicas_plataformas']
                t1, t2, t3, t4, t5, t6 = st.tabs(["Gupy", "inHire", "LinkedIn", "Indeed", "Catho", "Infojobs"])
                with t1: st.info(dicas.get('Gupy', 'Sem dicas específicas.'))
                with t2: st.info(dicas.get('inHire', 'Sem dicas específicas.'))
                with t3: st.info(dicas.get('LinkedIn', 'Sem dicas específicas.'))
                with t4: st.info(dicas.get('Indeed', 'Sem dicas específicas.'))
                with t5: st.info(dicas.get('Catho', 'Sem dicas específicas.'))
                with t6: st.info(dicas.get('Infojobs', 'Sem dicas específicas.'))

                st.divider()

                # Area de Download
                st.header("💾 Baixar Currículo Otimizado")
                st.markdown("Seu currículo foi reescrito (sem inventar dados) para maximizar a leitura linear dos robôs.")

                cv_md = resultado_json['cv_adaptado_markdown']

                col_btn1, col_btn2, col_btn3 = st.columns(3)

                with col_btn1:
                    st.download_button(
                        label="Baixar em Markdown (.md)",
                        data=cv_md,
                        file_name="cv_otimizado.md",
                        mime="text/markdown"
                    )

                with col_btn2:
                    pdf_bytes = gerar_pdf(cv_md)
                    st.download_button(
                        label="Baixar em PDF (.pdf)",
                        data=pdf_bytes,
                        file_name="cv_otimizado.pdf",
                        mime="application/pdf"
                    )

                with col_btn3:
                    docx_bytes = gerar_docx(cv_md)
                    st.download_button(
                        label="Baixar em Word (.docx)",
                        data=docx_bytes,
                        file_name="cv_otimizado.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    )

                with st.expander("Visualizar CV Otimizado"):
                    st.markdown(cv_md)

            except Exception as e:
                st.error(f"Erro na API ou no processamento: {e}")
    else:
        st.warning("Preencha a vaga e envie o currículo.")